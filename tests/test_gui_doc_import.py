"""Word/Excel 导入单测（用代码生成临时 docx/xlsx，不依赖 Office）。"""

from pathlib import Path

import pytest

pytest.importorskip("docx")
pytest.importorskip("openpyxl")

from megapro.gui.layout.doc_import import extract_docx, extract_xlsx


def _make_docx(path: Path):
    import docx

    d = docx.Document()
    d.add_paragraph("你好 World")
    d.add_paragraph("第二段")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A1"
    t.cell(0, 1).text = "B1"
    t.cell(1, 0).text = "A2"
    t.cell(1, 1).text = "B2"
    d.save(str(path))
    return path


def _make_xlsx(path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "姓名"
    ws["B1"] = "数量"
    ws["A2"] = "笔"
    ws["B2"] = 12
    wb.save(str(path))
    return path


def test_extract_docx_paragraphs_and_table(tmp_path):
    p = _make_docx(tmp_path / "t.docx")
    items = extract_docx(p, mode="outline", size_mm=5.0)
    # 2 段文字 + 表格（网格 + 4 格文字）= 至少 3 个 item
    assert len(items) >= 3
    names = [it.name for it in items]
    assert any("段落" in n for n in names)
    assert any("表" in n for n in names)
    # 所有路径坐标应在床面内（fit 后）
    for it in items:
        for pl in it.paths:
            for x, y in pl:
                assert 0 <= x <= 210 and 0 <= y <= 210


def test_extract_xlsx_grid_and_cells(tmp_path):
    p = _make_xlsx(tmp_path / "t.xlsx")
    items = extract_xlsx(p, mode="outline", size_mm=5.0)
    names = [it.name for it in items]
    assert any("网格" in n for n in names)
    assert any("格:" in n for n in names)
    for it in items:
        for pl in it.paths:
            for x, y in pl:
                assert 0 <= x <= 210 and 0 <= y <= 210


def test_extract_xlsx_fits_bed(tmp_path):
    # 表格应 fit 到 210×210（用小表保持测试快）
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 5):
        for c in range(1, 4):
            ws.cell(row=r, column=c, value=f"{r},{c}")
    p = tmp_path / "big.xlsx"
    wb.save(str(p))
    items = extract_xlsx(p, mode="single", size_mm=4.0)
    xs = [x for it in items for pl in it.paths for x, y in pl]
    ys = [y for it in items for pl in it.paths for x, y in pl]
    assert max(xs) <= 210 and max(ys) <= 210


# --- 阶段 2：行序方向断言（Word/Excel 正立，§7/§2.3 #7） ---------------------

def _yc(items, name_prefix):
    """具名图元的页面几何 y 中心（transformed_paths，纸面 y-up）。"""
    it = next(i for i in items if i.name.startswith(name_prefix))
    ys = [y for p in it.transformed_paths() for _, y in p]
    return sum(ys) / len(ys)


def _xc(items, name_prefix):
    """具名图元的页面几何 x 中心。"""
    it = next(i for i in items if i.name.startswith(name_prefix))
    xs = [x for p in it.transformed_paths() for x, _ in p]
    return sum(xs) / len(xs)


def test_extract_docx_upright_row_order(tmp_path):
    """Word 段落/表格行序正立：第一段/第一行在纸面 y-up 的上方。"""
    p = _make_docx(tmp_path / "t.docx")
    items = extract_docx(p, mode="outline", size_mm=5.0)
    # 段落：『你好 World』在『第二段』之上
    assert _yc(items, "段落:你好") > _yc(items, "段落:第二")
    # 表格：第一行（A1/B1）在第二行（A2/B2）之上
    assert _yc(items, "格:A1") > _yc(items, "格:A2")
    assert _yc(items, "格:B1") > _yc(items, "格:B2")
    # 相对布局保持（组契约：列序不变）
    assert _xc(items, "格:A1") < _xc(items, "格:B1")
    assert _xc(items, "格:A2") < _xc(items, "格:B2")


def test_extract_xlsx_upright_row_order(tmp_path):
    """Excel 行序正立 + 相对布局保持。"""
    p = _make_xlsx(tmp_path / "t.xlsx")
    items = extract_xlsx(p, mode="outline", size_mm=5.0)
    assert _yc(items, "格:姓名") > _yc(items, "格:笔")
    assert _yc(items, "格:数量") > _yc(items, "格:12")
    assert _xc(items, "格:姓名") < _xc(items, "格:数量")
    assert _xc(items, "格:笔") < _xc(items, "格:12")


def test_extract_docx_group_layout_preserved(tmp_path):
    """组导入契约：整组一次归位 —— 段落间距（相对布局）入库后保持非零。"""
    p = _make_docx(tmp_path / "t.docx")
    items = extract_docx(p, mode="outline", size_mm=5.0)
    gap = _yc(items, "段落:你好") - _yc(items, "段落:第二")
    assert gap > 1.0  # 不塌到同一点（严禁逐 Item 归位）
    # 全部坐标仍在床内（fit + 组归位后）
    for it in items:
        for pl in it.transformed_paths():
            for x, y in pl:
                assert 0 <= x <= 210 and 0 <= y <= 210
